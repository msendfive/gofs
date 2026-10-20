import os
import re
import html
import base64
from io import BytesIO
from datetime import datetime, date, timedelta, timezone
from email.utils import parsedate_to_datetime
from xml.sax.saxutils import escape

import pandas as pd
import streamlit as st
from google.auth.transport.requests import Request
from google.oauth2.credentials import Credentials
from google_auth_oauthlib.flow import Flow
from googleapiclient.discovery import build

from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.enums import TA_LEFT, TA_RIGHT
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle

from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

st.set_page_config(page_title="Gojek Report Generator", page_icon="📄", layout="wide")

handle_oauth_callback()

SCOPES = ['https://www.googleapis.com/auth/gmail.readonly']
WIB = timezone(timedelta(hours=7))

DARK = colors.HexColor("#1A1A1A")
GREEN = colors.HexColor("#00AA13")
GREY_BOX = colors.HexColor("#F2F2F2")
GREY_LINE = colors.HexColor("#CFCFCF")
DATE_FMT = "%d/%m/%Y %I:%M:%S %p"   # contoh: 13/10/2026 11:41:13 AM

COLUMNS = ['Tanggal', 'No. Transaksi', 'Layanan', 'Dari', 'Alamat Dari',
           'Tujuan', 'Alamat Tujuan', 'Metode Bayar', 'Total']
TEXT_COLS = [c for c in COLUMNS if c != 'Total']


def rp(n):
    return "Rp" + f"{int(n):,}".replace(",", ".")


# ------------------------------------------------------------------
# GMAIL
# ------------------------------------------------------------------
os.environ["OAUTHLIB_RELAX_TOKEN_SCOPE"] = "1"


def _make_flow():
    g = st.secrets["google"]
    cfg = {"web": {
        "client_id": g["client_id"],
        "client_secret": g["client_secret"],
        "auth_uri": "https://accounts.google.com/o/oauth2/auth",
        "token_uri": "https://oauth2.googleapis.com/token",
        "redirect_uris": [g["redirect_uri"]],
    }}
    # PKCE dimatikan karena sesi Streamlit baru dibuat setelah redirect,
    # sehingga code_verifier tidak akan tersimpan.
    return Flow.from_client_config(
        cfg, scopes=SCOPES, redirect_uri=g["redirect_uri"],
        autogenerate_code_verifier=False,
    )


def handle_oauth_callback():
    """Tukar ?code=... dari Google menjadi kredensial (dipanggil di awal script)."""
    code = st.query_params.get("code")
    if code and "creds" not in st.session_state:
        try:
            flow = _make_flow()
            flow.fetch_token(code=code)
            st.session_state["creds"] = flow.credentials
        except Exception as e:
            st.error(f"Login Google gagal: {e}")
        st.query_params.clear()
        st.rerun()


def get_gmail_service():
    creds = st.session_state.get("creds")
    if creds and creds.expired and creds.refresh_token:
        creds.refresh(Request())
        st.session_state["creds"] = creds
    if not creds or not creds.valid:
        auth_url, _ = _make_flow().authorization_url(
            access_type="offline", prompt="consent", include_granted_scopes="true"
        )
        st.warning("Anda belum login ke Google.")
        st.markdown(
            f'<a href="{auth_url}" target="_self">🔐 <b>Login dengan Google</b></a>',
            unsafe_allow_html=True,
        )
        return None
    return build("gmail", "v1", credentials=creds)


def extract_body(payload):
    """Ambil teks email (text/plain, fallback HTML) dengan baris baru tetap dipertahankan."""
    def walk(part, mime):
        if part.get('mimeType') == mime:
            data = part.get('body', {}).get('data')
            if data:
                return base64.urlsafe_b64decode(data).decode('utf-8', errors='ignore')
        for sub in part.get('parts', []) or []:
            res = walk(sub, mime)
            if res:
                return res
        return ""

    text = walk(payload, 'text/plain')
    if not text:
        raw = walk(payload, 'text/html')
        raw = re.sub(r'(?i)<(br|/p|/tr|/td|/th|/div|/li|/h\d)[^>]*>', '\n', raw)
        raw = re.sub(r'(?is)<(script|style).*?</\1>', ' ', raw)
        text = html.unescape(re.sub(r'<[^>]+>', ' ', raw))
    lines = [re.sub(r'[ \t\u00a0]+', ' ', ln).strip() for ln in text.splitlines()]
    return "\n".join(ln for ln in lines if ln)


# ------------------------------------------------------------------
# PARSING EMAIL GOJEK (best-effort, hasil bisa dikoreksi di tabel editor)
# ------------------------------------------------------------------
def to_int(s):
    s = s.replace('.', '').split(',')[0]
    return int(s) if s.isdigit() else 0


def find_after(lines, labels):
    """Cari baris berlabel (mis. 'Dari: X' atau 'Dari' lalu baris berikutnya 'X')."""
    pat = re.compile(r'^(?:%s)\s*:?\s*(.*)$' % '|'.join(labels), re.I)
    for i, ln in enumerate(lines):
        m = pat.match(ln)
        if m:
            val = m.group(1).strip()
            if val:
                return val
            if i + 1 < len(lines):
                return lines[i + 1]
    return ""


def after_marker(lines, pattern):
    """Ambil 2 baris setelah baris penanda (nama tempat, lalu alamatnya).
    Contoh: 'Diantarkan jam 18:30 dari' -> 'Teguk, Cibaduyut' -> 'Jl. Terusan Cibaduyut No. 104...'"""
    pat = re.compile(pattern, re.I)
    for i, ln in enumerate(lines):
        m = pat.match(ln)
        if m:
            inline = ln[m.end():].strip()
            rest = ([inline] if inline else []) + lines[i + 1:] + ["", ""]
            return rest[0], rest[1]
    return "", ""


def detect_layanan(subject, body, order_id):
    # ID pesanan paling pasti: F-xxxx = GoFood, GK-xx-xxxx = GoSend
    if order_id.startswith('F-'):
        return 'GoFood'
    if order_id.startswith('GK-'):
        return 'GoSend'
    keys_map = (('GoFood', ['gofood', 'go-food']), ('GoSend', ['gosend', 'go-send']))
    s = subject.lower()
    for name, keys in keys_map:
        if any(k in s for k in keys):
            return name
    b = body.lower()
    pos = {}
    for name, keys in keys_map:
        idx = [b.find(k) for k in keys if k in b]
        if idx:
            pos[name] = min(idx)
    return min(pos, key=pos.get) if pos else None


def find_total(body):
    for label in (r'Total\s+dibayar', r'Total\s+pembayaran'):
        m = re.search(label + r'\D{0,40}?Rp\s?([\d.,]+)', body, re.I)
        if m:
            return to_int(m.group(1))
    totals = re.findall(r'(?i)(?<![a-z])total\D{0,40}?Rp\s?([\d.,]+)', body)
    if totals:
        return to_int(totals[-1])
    allrp = [to_int(x) for x in re.findall(r'Rp\s?([\d.,]+)', body)]
    return max(allrp) if allrp else 0


def find_payment(lines, body):
    """'Bayar pakai GoPay' + 'Bayar pakai Cash' -> 'GoPay + Cash'."""
    methods = []
    for ln in lines:
        m = re.match(r'(?i)bayar\s+pakai\s+(.+?)\s*(?:Rp.*)?$', ln)
        if m and m.group(1).strip() and m.group(1).strip() not in methods:
            methods.append(m.group(1).strip())
    if methods:
        return " + ".join(methods)
    m = re.search(r'(?i)\b(GoPay Later|GoPay|Tunai|Cash|OVO|QRIS|ShopeePay|Kartu Kredit|Kartu Debit)\b', body)
    return m.group(1) if m else ""


def parse_gojek_email(msg_id, headers, body):
    h = {x.get('name'): x.get('value') for x in headers}
    subject = h.get('Subject', '')
    lines = body.splitlines()

    m = re.search(r'(?i)ID\s+pesanan\s*:?\s*([A-Z]{1,3}-[\w-]+)', body)
    if not m:
        m = re.search(r'\b(F-\d{6,}|GK-\d+-\d+)\b', f"{subject}\n{body}")
    order_id = m.group(1) if m else ""

    layanan = detect_layanan(subject, body, order_id)
    if not layanan:
        return None  # bukan GoFood / GoSend (mis. promo atau GoRide)

    try:
        dt = parsedate_to_datetime(h.get('Date', '')).astimezone(WIB).replace(tzinfo=None)
    except Exception:
        dt = datetime.min

    dari, alamat_dari = after_marker(lines, r'(?:Diantarkan|Dijemput|Diambil)\s+jam\s+[\d:.]+\s+(?:dari|di)\b')
    tujuan, alamat_tujuan = after_marker(lines, r'Sampai\s+jam\s+[\d:.]+\s+di\b')

    # cadangan jika format email berbeda
    if not dari:
        mm = re.search(r'(?i)pesan makan dari (.+)', body)
        dari = mm.group(1).strip() if mm else find_after(lines, ['Restoran', 'Merchant', 'Dari', 'Pickup', 'Lokasi jemput'])
    if not tujuan:
        tujuan = find_after(lines, ['Tujuan', 'Diantar ke', 'Alamat pengantaran', 'Lokasi antar', 'Alamat tujuan'])

    return {
        'Tanggal': dt.strftime(DATE_FMT) if dt != datetime.min else h.get('Date', ''),
        'No. Transaksi': order_id or f"GJ-{msg_id[-8:].upper()}",
        'Layanan': layanan,
        'Dari': dari,
        'Alamat Dari': alamat_dari,
        'Tujuan': tujuan,
        'Alamat Tujuan': alamat_tujuan,
        'Metode Bayar': find_payment(lines, body),
        'Total': find_total(body),
        '_sort': dt,
    }


SAMPLE_ROWS = [
    ('13/10/2026 11:41:13 AM', 'F-3389892933', 'GoFood', 'Gepuk Ny. Yong, Gandapura',
     'Jl. Gandapura 39 RT 04 RW 04, Kel. Merdeka, Kec. Sumur Bandung, Kode Pos 40113',
     'Jl. Sudirman No.823', 'Jl. Jenderal Sudirman No.823, Cijerah, Kec. Bandung Kulon, Kota Bandung, Jawa Barat 40213, Indonesia', 'GoPay', 178050),
    ('12/10/2026 02:30:10 PM', 'F-3389011146', 'GoFood', 'Kurokoffee, Burwood Sports Club',
     '4G3V+VJP, Jl. HMS Mintareja Sarjana Hukum, Baros, Kec. Cimahi Tengah, Kota Cimahi, Jawa Barat 40521, Indonesia',
     'Jl. Sudirman No.823', 'Jl. Jenderal Sudirman No.823, Cijerah, Kec. Bandung Kulon, Kota Bandung, Jawa Barat 40213, Indonesia', 'GoPay', 51700),
    ('10/10/2026 11:25:34 AM', 'F-3386272878', 'GoFood', 'Tabulla Rasa, Dapur Oleh Pondok Ikan Cianjur, Bandung Kulon',
     'Jl. Soekarno Hatta No. 88, Bandung Kulon, Bandung',
     'Jl. Sudirman No.823', 'Jl. Jenderal Sudirman No.823, Cijerah, Kec. Bandung Kulon, Kota Bandung, Jawa Barat 40213, Indonesia', 'GoPay', 258000),
    ('09/10/2026 03:29:27 PM', 'GK-91-341407399', 'GoSend', 'Jl. Jenderal Sudirman No.823',
     'Jl. Jenderal Sudirman No.823, Cijerah, Kec. Bandung Kulon, Kota Bandung, Jawa Barat 40213, Indonesia',
     'Kopo Mas Regency', 'Kopo Mas Regency, Blk. B No.25, Margasuka, Kec. Babakan Ciparay, Kabupaten Bandung, Jawa Barat 40225, Indonesia', 'GoPay', 30000),
    ('08/10/2026 02:42:31 PM', 'F-3384238858', 'GoFood', 'Kurokoffee, Burwood Sports Club',
     '4G3V+VJP, Jl. HMS Mintareja Sarjana Hukum, Baros, Kec. Cimahi Tengah, Kota Cimahi, Jawa Barat 40521, Indonesia',
     'Jl. Sudirman No.823', 'Jl. Jenderal Sudirman No.823, Cijerah, Kec. Bandung Kulon, Kota Bandung, Jawa Barat 40213, Indonesia', 'GoPay', 51700),
    ('08/10/2026 02:30:19 PM', 'GK-91-341061524', 'GoSend', 'Jl. Jenderal Sudirman No.823',
     'Jl. Jenderal Sudirman No.823, Cijerah, Kec. Bandung Kulon, Kota Bandung, Jawa Barat 40213, Indonesia',
     'Blk. A3 No.16', 'Blk. A3 No.16, Cibaduyut Kidul, Kec. Bojongloa Kidul, Kota Bandung, Jawa Barat 40227, Indonesia', 'GoPay', 10000),
    ('08/10/2026 01:38:20 PM', 'GK-91-341041260', 'GoSend', 'Jl. Jenderal Sudirman No.823',
     'Jl. Jenderal Sudirman No.823, Cijerah, Kec. Bandung Kulon, Kota Bandung, Jawa Barat 40213, Indonesia',
     'Jl. Dian Sakti No.11', 'Jl. Dian Sakti No.11, Babakan, Kec. Babakan Ciparay, Kota Bandung, Jawa Barat 40222, Indonesia', 'GoPay', 10000),
    ('08/10/2026 12:25:04 PM', 'F-3384081548', 'GoFood', 'Gormeteria, Cicendo',
     'Jl. Pasirkaliki No. 176, Cicendo, Bandung',
     'Jl. Sudirman No.823', 'Jl. Jenderal Sudirman No.823, Cijerah, Kec. Bandung Kulon, Kota Bandung, Jawa Barat 40213, Indonesia', 'GoPay', 83400),
]


def sample_dataframe():
    return pd.DataFrame(SAMPLE_ROWS, columns=COLUMNS)


# ------------------------------------------------------------------
# PDF  (gaya laporan transaksi Gojek)
# ------------------------------------------------------------------
def _draw_page(canvas, doc, label):
    w, h = A4
    canvas.saveState()
    canvas.setFillColor(DARK)
    canvas.roundRect(22, h - 68, w - 44, 46, 10, fill=1, stroke=0)
    # logo sederhana: lingkaran putih + teks gojek
    canvas.setStrokeColor(colors.white)
    canvas.setLineWidth(2)
    canvas.circle(46, h - 45, 8, fill=0, stroke=1)
    canvas.setFillColor(colors.white)
    canvas.circle(46, h - 45, 3, fill=1, stroke=0)
    canvas.setFont("Helvetica-Bold", 15)
    canvas.drawString(60, h - 50, "gojek")
    canvas.setFont("Helvetica", 8)
    canvas.drawRightString(w - 38, h - 48, label)
    canvas.setFillColor(colors.HexColor("#777777"))
    canvas.setFont("Helvetica", 8)
    canvas.drawRightString(w - 30, 18, f"Halaman {doc.page}")
    canvas.restoreState()


def build_pdf(df, nama, email, telepon, layanan_label):
    buf = BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, rightMargin=30, leftMargin=30,
                            topMargin=88, bottomMargin=34)
    el = []

    name = ParagraphStyle('name', fontName='Helvetica-Bold', fontSize=12, leading=15)
    small = ParagraphStyle('small', fontName='Helvetica', fontSize=8.5, leading=11)
    big = ParagraphStyle('big', parent=small, fontName='Helvetica-Bold', fontSize=11, leading=14)
    cell = ParagraphStyle('cell', fontName='Helvetica', fontSize=7, leading=9, alignment=TA_LEFT)
    cellb = ParagraphStyle('cellb', parent=cell, fontName='Helvetica-Bold')
    cellr = ParagraphStyle('cellr', parent=cell, alignment=TA_RIGHT)
    cellrb = ParagraphStyle('cellrb', parent=cellr, fontName='Helvetica-Bold', fontSize=8)
    head = ParagraphStyle('head', fontName='Helvetica-Bold', fontSize=8, leading=10)
    headr = ParagraphStyle('headr', parent=head, alignment=TA_RIGHT)

    dates = pd.to_datetime(df['Tanggal'], format=DATE_FMT, errors='coerce').dropna()
    periode = f"{dates.min():%d %B %Y} - {dates.max():%d %B %Y}" if len(dates) else "-"
    total_all = int(pd.to_numeric(df['Total'], errors='coerce').fillna(0).sum())

    left = [Paragraph(escape(nama or "-"), name), Spacer(1, 3)]
    if telepon:
        left.append(Paragraph(escape(telepon), small))
    left.append(Paragraph(escape(email or "-"), small))
    right = [Paragraph("Periode transaksi:", small), Paragraph(escape(periode), small), Spacer(1, 4),
             Paragraph("Total transaksi:", small), Paragraph(rp(total_all), big)]

    info = Table([[left, right]], colWidths=[250, 285])
    info.setStyle(TableStyle([
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('BACKGROUND', (0, 0), (-1, -1), GREY_BOX),
        ('ROUNDEDCORNERS', [8, 8, 8, 8]),
        ('LEFTPADDING', (0, 0), (-1, -1), 12),
        ('TOPPADDING', (0, 0), (-1, -1), 10),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 10),
    ]))
    el += [info, Spacer(1, 12),
           Paragraph(f"Bukti satuan dari semua transaksi {escape(layanan_label)} berikut bisa didownload di aplikasi.", small),
           Spacer(1, 6)]

    def place(name_, addr):
        parts = [Paragraph(escape(str(name_)), cellb)]
        if str(addr).strip():
            parts += [Spacer(1, 2), Paragraph(escape(str(addr)), cell)]
        return parts

    rows = [[Paragraph('Tanggal', head), Paragraph('No.<br/>transaksi', head), Paragraph('Layanan', head),
             Paragraph('Dari', head), Paragraph('Tujuan', head), Paragraph('Metode<br/>bayar', head),
             Paragraph('Total<br/>dibayar', headr)]]
    for _, r in df.iterrows():
        rows.append([
            Paragraph(escape(str(r['Tanggal'])).replace(' ', '<br/>', 1), cell),
            Paragraph(escape(str(r['No. Transaksi'])), cell),
            Paragraph(escape(str(r['Layanan'])), cell),
            place(r['Dari'], r['Alamat Dari']),
            place(r['Tujuan'], r['Alamat Tujuan']),
            Paragraph(escape(str(r['Metode Bayar'])), cell),
            Paragraph(rp(r['Total']), cellr),
        ])
    rows.append([Paragraph('Total', cellrb), '', '', '', '', '', Paragraph(rp(total_all), cellrb)])
    last = len(rows) - 1

    t = Table(rows, colWidths=[58, 58, 46, 135, 125, 38, 75], repeatRows=1)
    t.setStyle(TableStyle([
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('LINEABOVE', (0, 0), (-1, 0), 0.8, DARK),
        ('LINEBELOW', (0, 0), (-1, 0), 0.8, DARK),
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#FAFAFA")),
        ('GRID', (0, 1), (-1, -1), 0.5, GREY_LINE),
        ('BOX', (0, 0), (-1, -1), 0.5, GREY_LINE),
        ('SPAN', (0, last), (5, last)),
        ('BACKGROUND', (0, last), (-1, last), colors.HexColor("#FAFAFA")),
        ('LEFTPADDING', (0, 0), (-1, -1), 4),
        ('RIGHTPADDING', (0, 0), (-1, -1), 4),
        ('TOPPADDING', (0, 0), (-1, -1), 5),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 7),
    ]))
    el.append(t)
    doc.build(el, onFirstPage=lambda c, d: _draw_page(c, d, f"Laporan {layanan_label}"),
              onLaterPages=lambda c, d: _draw_page(c, d, f"Laporan {layanan_label}"))
    return buf.getvalue()


# ------------------------------------------------------------------
# EXCEL
# ------------------------------------------------------------------
def build_excel(df, nama, email, layanan_label):
    out = df[COLUMNS].copy()
    out['Total'] = pd.to_numeric(out['Total'], errors='coerce').fillna(0).astype(int)

    buf = BytesIO()
    with pd.ExcelWriter(buf, engine='openpyxl') as writer:
        out.to_excel(writer, index=False, sheet_name=layanan_label[:30], startrow=4)
        ws = writer.sheets[layanan_label[:30]]

        ws['A1'] = f"LAPORAN TRANSAKSI {layanan_label.upper()}"
        ws['A1'].font = Font(bold=True, size=14)
        ws['A2'] = f"{nama or '-'}  |  {email or '-'}"
        ws['A3'] = f"Total transaksi: {len(out)}   |   Dibuat: {datetime.now():%d/%m/%Y %H:%M}"

        thin = Side(style='thin', color='CFCFCF')
        border = Border(left=thin, right=thin, top=thin, bottom=thin)
        hdr = 5
        ncol = len(COLUMNS)
        for c in range(1, ncol + 1):
            cell = ws.cell(row=hdr, column=c)
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='1A1A1A')
            cell.alignment = Alignment(horizontal='center', vertical='center')
            cell.border = border

        first, last = hdr + 1, ws.max_row
        for row in ws.iter_rows(min_row=first, max_row=last, max_col=ncol):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical='top')
                cell.border = border
            row[ncol - 1].number_format = '"Rp"#,##0'
            row[ncol - 1].alignment = Alignment(horizontal='right', vertical='top')

        tr = last + 1
        ws.cell(row=tr, column=ncol - 1, value="TOTAL").font = Font(bold=True)
        ws.cell(row=tr, column=ncol - 1).alignment = Alignment(horizontal='right')
        tc = ws.cell(row=tr, column=ncol, value=f"=SUM({get_column_letter(ncol)}{first}:{get_column_letter(ncol)}{last})")
        tc.font = Font(bold=True)
        tc.number_format = '"Rp"#,##0'

        for i, w in enumerate([22, 20, 10, 30, 45, 24, 45, 14, 14], start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = ws.cell(row=first, column=1)
        ws.auto_filter.ref = f"A{hdr}:{get_column_letter(ncol)}{last}"
    return buf.getvalue()


# ------------------------------------------------------------------
# UI
# ------------------------------------------------------------------
def render_section(label, key, icon):
    """Satu bagian terpisah (GoFood / GoSend): editor + download PDF & Excel."""
    base = st.session_state[f'df_{key}'].copy()
    if base.empty:
        st.info(f"Tidak ada transaksi {label} pada hasil tarikan.")
        return
    base.insert(0, 'Sertakan', True)

    st.caption("Klik sel untuk mengedit. Hilangkan centang **Sertakan** untuk mengecualikan baris. "
               "Pilih baris lalu tekan Delete untuk menghapus, atau tambah baris di bagian bawah tabel.")
    edited = st.data_editor(
        base,
        key=f"editor_{key}_{st.session_state.get('ver', 0)}",
        num_rows="dynamic", hide_index=True, use_container_width=True,
        column_config={
            'Sertakan': st.column_config.CheckboxColumn("Sertakan", default=True, width="small"),
            'Layanan': st.column_config.TextColumn("Layanan", default=label, width="small"),
            'Alamat Dari': st.column_config.TextColumn("Alamat Dari", width="large"),
            'Alamat Tujuan': st.column_config.TextColumn("Alamat Tujuan", width="large"),
            'Total': st.column_config.NumberColumn("Total", format="Rp %d", min_value=0, step=1, default=0),
        },
    )

    final = edited.copy()
    final['Sertakan'] = final['Sertakan'].fillna(True).astype(bool)
    final = final[final['Sertakan']].drop(columns='Sertakan')
    final['Total'] = pd.to_numeric(final['Total'], errors='coerce').fillna(0).astype(int)
    final[TEXT_COLS] = final[TEXT_COLS].fillna("").astype(str)
    nonempty = final[TEXT_COLS].apply(lambda c: c.str.strip()).ne("").any(axis=1) | final['Total'].gt(0)
    final = final[nonempty].reset_index(drop=True)

    if final.empty:
        st.warning("Tidak ada baris yang dipilih untuk dimasukkan ke laporan.")
        return

    st.info(f"{len(final)} transaksi {label} • Total {rp(final['Total'].sum())}")
    c1, c2 = st.columns(2)
    c1.download_button(f"📄 Download PDF {label}",
                       build_pdf(final, st.session_state['nama'], st.session_state['email'],
                                 st.session_state['telepon'], label),
                       file_name=f"Laporan_{label}.pdf", mime="application/pdf", key=f"pdf_{key}")
    c2.download_button(f"📊 Download Excel {label}",
                       build_excel(final, st.session_state['nama'], st.session_state['email'], label),
                       file_name=f"Laporan_{label}.xlsx",
                       mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                       key=f"xls_{key}")


st.title("🚀 Gojek Email to PDF & Excel Report Generator")
st.markdown("Tarik email transaksi **Gojek** dari Gmail, pisahkan **GoFood** dan **GoSend**, "
            "periksa/edit datanya, lalu ekspor ke **PDF** atau **Excel**.")

today = date.today()
with st.sidebar:
    st.header("⚙️ Konfigurasi")
    if "creds" in st.session_state:
        st.success("✅ Sudah login ke Google")
        if st.button("Logout"):
            del st.session_state["creds"]
            st.rerun()
    else:
        st.info("Belum login. Klik 'Tarik Data Email' untuk mendapatkan link login Google.")

    st.subheader("Data Laporan")
    nama_in = st.text_input("Nama di laporan", value="GAP")
    telepon_in = st.text_input("No. telepon (opsional)", value="")
    rng = st.date_input("Periode transaksi", value=(today.replace(day=1), today))
    if isinstance(rng, (tuple, list)):
        start = rng[0]
        end = rng[1] if len(rng) > 1 else rng[0]
    else:
        start = end = rng
    sender_query = st.text_input("Filter pengirim (query Gmail)", value="from:invoicing.gojek.com",
                                 help="Email struk Gojek dikirim dari no-reply@invoicing.gojek.com. Ubah jika alamatmu berbeda.")
    max_results = st.slider("Jumlah email maksimal", 10, 200, 50)

if st.button("📥 Tarik Data Email", type="primary"):
    service = get_gmail_service()
    if service:
        with st.spinner("Sedang memproses email dari Gmail..."):
            try:
                email_acc = service.users().getProfile(userId='me').execute().get('emailAddress', '')
            except Exception:
                email_acc = ''

            query = f"{sender_query} after:{start:%Y/%m/%d} before:{end + timedelta(days=1):%Y/%m/%d}"
            results = service.users().messages().list(userId='me', q=query, maxResults=max_results).execute()
            messages = results.get('messages', [])

            rows, raw_samples = [], []
            for msg in messages:
                try:
                    txt = service.users().messages().get(userId='me', id=msg['id']).execute()
                    payload = txt.get('payload', {})
                    body = extract_body(payload)
                    parsed = parse_gojek_email(msg['id'], payload.get('headers', []), body)
                    if parsed:
                        rows.append(parsed)
                        if len(raw_samples) < 2:
                            raw_samples.append(body[:1500])
                except Exception as e:
                    print(f"Error: {e}")

            if rows:
                df = pd.DataFrame(rows).sort_values('_sort', ascending=False).drop(columns='_sort')
                st.session_state['sample_mode'] = False
            else:
                st.warning("Tidak ada email GoFood/GoSend ditemukan. Menampilkan data contoh untuk uji coba.")
                df = sample_dataframe()
                st.session_state['sample_mode'] = True

            df = df[COLUMNS].reset_index(drop=True)
            st.session_state['df_gofood'] = df[df['Layanan'] == 'GoFood'].reset_index(drop=True)
            st.session_state['df_gosend'] = df[df['Layanan'] == 'GoSend'].reset_index(drop=True)
            st.session_state['nama'] = nama_in
            st.session_state['telepon'] = telepon_in
            st.session_state['email'] = email_acc
            st.session_state['raw_samples'] = raw_samples
            st.session_state['ver'] = st.session_state.get('ver', 0) + 1

if 'df_gofood' in st.session_state:
    # nama & telepon bisa diubah di sidebar tanpa menarik ulang email
    st.session_state['nama'] = nama_in
    st.session_state['telepon'] = telepon_in

    tab_food, tab_send = st.tabs(["🍔 GoFood", "📦 GoSend"])
    with tab_food:
        render_section("GoFood", "gofood", "🍔")
    with tab_send:
        render_section("GoSend", "gosend", "📦")

    if st.session_state.get('raw_samples'):
        with st.expander("🔍 Lihat isi email mentah (untuk mengecek hasil parsing)"):
            for i, raw in enumerate(st.session_state['raw_samples'], 1):
                st.text(f"--- Email {i} ---\n{raw}")